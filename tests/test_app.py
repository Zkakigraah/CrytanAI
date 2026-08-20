"""
Tests for app/file_extraction.py and the newer parts of app/pipeline.py
(explain_result, analyze_text) -- all pure functions, no Streamlit
required to test them, which is exactly why that separation exists.
"""
import io

import pytest

from app.file_extraction import extract_text
from app.pipeline import analyze_text, explain_result, detect_and_solve
from ciphers.registry import get_cipher


# --------------------------------------------------------------- extraction

def test_extract_txt():
    result = extract_text("note.txt", b"Hello World")
    assert result == "Hello World"


def test_extract_md_strips_formatting():
    md = b"# Heading\n\nThis is **bold** and *italic* with a [link](http://x.com).\n\n- bullet one\n- bullet two"
    result = extract_text("note.md", md)
    assert "#" not in result
    assert "**" not in result
    assert "[" not in result and "](http" not in result
    assert "Heading" in result
    assert "bold" in result and "italic" in result
    assert "bullet one" in result


def test_extract_md_strips_table_separators():
    md = b"| A | B |\n|---|---|\n| 1 | 2 |"
    result = extract_text("table.md", md)
    assert "---" not in result
    assert "|" not in result


def test_extract_docx():
    from docx import Document

    doc = Document()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    buf = io.BytesIO()
    doc.save(buf)

    result = extract_text("doc.docx", buf.getvalue())
    assert "First paragraph." in result
    assert "Second paragraph." in result


def test_extract_pdf():
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 750, "PDF extraction test line.")
    c.save()

    result = extract_text("doc.pdf", buf.getvalue())
    assert "PDF extraction test line" in result


def test_extract_unsupported_extension_raises():
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text("file.xyz", b"data")


def test_extract_handles_latin1_fallback():
    # A byte sequence that isn't valid UTF-8 but is valid Latin-1
    raw = "café".encode("latin-1")
    result = extract_text("note.txt", raw)
    assert "caf" in result  # doesn't crash, recovers what it can


# --------------------------------------------------------------- analyze_text

LONG_ENGLISH = (
    "MACHINE LEARNING HAS TRANSFORMED THE FIELD OF CRYPTANALYSIS BY ALLOWING "
    "COMPUTERS TO AUTOMATICALLY DETECT PATTERNS IN ENCRYPTED TEXT."
)


def test_analyze_text_on_plain_english():
    analysis = analyze_text(LONG_ENGLISH)
    assert analysis.index_of_coincidence > 0.055  # close to expected English ~0.067
    assert analysis.word_ratio > 0.5  # most tokens are real words -- not stricter than
    # this, since LONG_ENGLISH intentionally includes technical vocabulary
    # ("CRYPTANALYSIS") that a general-purpose dictionary may reasonably lack
    assert analysis.digit_ratio == 0.0
    assert sum(analysis.observed_letter_freq.values()) == pytest.approx(1.0, abs=1e-6)
    assert set(analysis.observed_letter_freq.keys()) == set(analysis.expected_letter_freq.keys())


def test_analyze_text_flags_polybius_by_digit_ratio():
    spec = get_cipher("polybius")
    ciphertext = spec.encrypt(LONG_ENGLISH, "")
    analysis = analyze_text(ciphertext)
    assert analysis.digit_ratio > 0.8


def test_analyze_text_flags_playfair_by_token_length():
    from ml.features import extract_features, FEATURE_NAMES

    spec = get_cipher("playfair")
    ciphertext = spec.encrypt(LONG_ENGLISH, "EXAMPLE")
    features = dict(zip(FEATURE_NAMES, extract_features(ciphertext)))
    assert features["avg_token_length"] == pytest.approx(2.0, abs=0.1)
    assert features["token_length_variance"] < 0.1


def test_analyze_text_handles_empty_input():
    analysis = analyze_text("")
    assert analysis.length == 0
    assert analysis.word_ratio == 0.0


# --------------------------------------------------------------- explain_result

def test_explain_result_mentions_classifier_probability_and_method():
    spec = get_cipher("caesar")
    ciphertext = spec.encrypt(LONG_ENGLISH, spec.random_key())
    results = detect_and_solve(ciphertext, top_k=1)
    steps = explain_result(results[0])

    assert len(steps) >= 3
    assert any("probability" in s.lower() for s in steps)
    assert any("confidence" in s.lower() for s in steps)


def test_explain_result_handles_hill_suspected_case():
    from app.pipeline import DetectAndSolveResult

    hill_result = DetectAndSolveResult(
        cipher_slug="hill", cipher_display_name="Hill Cipher", cipher_family="polygraphic",
        key=None, plaintext="", score=float("-inf"), score_breakdown=None,
        confidence="n/a", method="none", classifier_probability=0.3,
        hill_suspected_but_needs_crib=True,
    )
    # Should not raise, even with score_breakdown=None and an empty plaintext.
    steps = explain_result(hill_result)
    assert len(steps) >= 1
